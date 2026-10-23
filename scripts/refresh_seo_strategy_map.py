#!/usr/bin/env python3
"""
scripts/refresh_seo_strategy_map.py

Reconciles and refreshes the SEO URL Strategy Map and related sheets in the Audiment SEO Database workbook:
1. Reconciles published URL (/blog/best-inspection-checklist-apps) with legacy longer URL.
2. Reconciles published URL (/blog/inspection-checklist-software-features) with legacy planning URL (/blog/inspection-checklist-software-app-features):
   - Sets /blog/inspection-checklist-software-features as the canonical published entry for "Inspection Checklist Software & App Features".
   - Retains C10 cluster mapping, inspection-management pillar, hubs, supporting links, and relationships.
   - Marks /blog/inspection-checklist-software-app-features as obsolete / never published for traceability.
3. Updates Internal Linking references (Rows 60-67, 170, 187, 202) to point to the canonical published URL while preserving traceability.
4. Updates Content Clusters (Row 9) Proposed URL to the published canonical URL.
5. Updates Question Map (22 rows) Proposed URL to the published canonical URL.
6. Dynamically computes the true rolling 28-day window from Page Performance (max_date - 27 days to max_date)
   and refreshes GSC metrics across all mapped URLs.
7. Updates the Last Metrics Refresh timestamp.
"""

import sys
import os
from datetime import datetime, timedelta
import openpyxl

def normalize_url(url: str) -> str:
    if not url:
        return ""
    u = str(url).strip()
    if u.startswith("http://"):
        u = "https://" + u[7:]
    if u.endswith("/") and u not in ("https://audiment.in/", "http://audiment.in/"):
        u = u.rstrip("/")
    elif u == "https://audiment.in":
        u = "https://audiment.in/"
    return u

def reconcile_and_refresh_workbook(workbook_path: str):
    if not os.path.exists(workbook_path):
        print(f"Error: Workbook not found at {workbook_path}")
        return False

    print(f"\n==========================================")
    print(f"Loading workbook: {workbook_path}")
    print(f"==========================================")
    wb = openpyxl.load_workbook(workbook_path, data_only=False)

    required_sheets = ['SEO URL Strategy Map', 'Page Performance']
    for s in required_sheets:
        if s not in wb.sheetnames:
            print(f"Required sheet '{s}' missing from workbook.")
            return False

    ws_sm = wb['SEO URL Strategy Map']
    ws_pp = wb['Page Performance']

    # 1. Determine Date Range in Page Performance for 28-day rolling window
    pp_dates = []
    for r in range(2, ws_pp.max_row + 1):
        v = ws_pp.cell(row=r, column=1).value
        if v:
            if isinstance(v, datetime):
                pp_dates.append(v.date())
            else:
                try:
                    pp_dates.append(datetime.strptime(str(v)[:10], '%Y-%m-%d').date())
                except Exception:
                    pass

    if not pp_dates:
        print("No valid dates found in Page Performance.")
        return False

    max_date = max(pp_dates)
    start_date = max_date - timedelta(days=27)  # 28 days inclusive
    print(f"Page Performance max date: {max_date}")
    print(f"Rolling 28-day window: {start_date} to {max_date} (28 days inclusive)")

    # 2. Aggregate 28-day GSC metrics per URL from Page Performance
    url_metrics = {}
    for r in range(2, ws_pp.max_row + 1):
        raw_url = ws_pp.cell(row=r, column=2).value
        if not raw_url:
            continue
        norm_url = normalize_url(raw_url)
        raw_dt = ws_pp.cell(row=r, column=1).value
        dt = raw_dt.date() if isinstance(raw_dt, datetime) else datetime.strptime(str(raw_dt)[:10], '%Y-%m-%d').date()

        if start_date <= dt <= max_date:
            clicks = float(ws_pp.cell(row=r, column=3).value or 0)
            impressions = float(ws_pp.cell(row=r, column=4).value or 0)
            position = float(ws_pp.cell(row=r, column=6).value or 0)

            if norm_url not in url_metrics:
                url_metrics[norm_url] = {"clicks": 0.0, "impressions": 0.0, "w_pos": 0.0}

            url_metrics[norm_url]["clicks"] += clicks
            url_metrics[norm_url]["impressions"] += impressions
            url_metrics[norm_url]["w_pos"] += position * impressions

    for u, data in url_metrics.items():
        imp = data["impressions"]
        c = data["clicks"]
        data["ctr"] = (c / imp) if imp > 0 else 0.0
        data["avg_pos"] = (data["w_pos"] / imp) if imp > 0 else 0.0

    # 3. Reconciliation 1: Best Inspection Checklist Apps
    APPS_LONGER_URL = "https://audiment.in/blog/best-inspection-checklist-apps-in-2026-reviewed-for-multi-location-teams"
    APPS_PUBLISHED_URL = "https://audiment.in/blog/best-inspection-checklist-apps"

    row_apps_longer = None
    row_apps_pub = None

    # Reconciliation 2: Inspection Checklist Software Features
    FEAT_OBSOLETE_URL = "https://audiment.in/blog/inspection-checklist-software-app-features"
    FEAT_PUBLISHED_URL = "https://audiment.in/blog/inspection-checklist-software-features"

    row_feat_obsolete = None
    row_feat_pub = None

    for r in range(2, ws_sm.max_row + 1):
        u = normalize_url(ws_sm.cell(row=r, column=1).value)
        if u == APPS_LONGER_URL:
            row_apps_longer = r
        elif u == APPS_PUBLISHED_URL:
            row_apps_pub = r
        elif u == FEAT_OBSOLETE_URL:
            row_feat_obsolete = r
        elif u == FEAT_PUBLISHED_URL:
            row_feat_pub = r

    print(f"Apps rows -> Longer/Alias: {row_apps_longer}, Published: {row_apps_pub}")
    print(f"Features rows -> Obsolete planning: {row_feat_obsolete}, Published: {row_feat_pub}")

    # Reconcile Apps rows
    if row_apps_longer and row_apps_pub:
        ws_sm.cell(row=row_apps_pub, column=2, value="Existing published page")
        ws_sm.cell(row=row_apps_pub, column=3, value=ws_sm.cell(row=row_apps_longer, column=3).value or "Best Inspection Checklist Apps in 2026 (Reviewed for Multi-Location Teams)")
        ws_sm.cell(row=row_apps_pub, column=4, value="Inspection-app commercial hub")
        ws_sm.cell(row=row_apps_pub, column=5, value=ws_sm.cell(row=row_apps_longer, column=5).value or "C04 | C10")
        ws_sm.cell(row=row_apps_pub, column=6, value=ws_sm.cell(row=row_apps_longer, column=6).value or "How to Evaluate Audit Management Software | Inspection Checklist Software & App Features")
        ws_sm.cell(row=row_apps_pub, column=7, value=ws_sm.cell(row=row_apps_longer, column=7).value or "Audit Management | Inspection Management")
        ws_sm.cell(row=row_apps_pub, column=8, value=ws_sm.cell(row=row_apps_longer, column=8).value or "High")
        ws_sm.cell(row=row_apps_pub, column=9, value=ws_sm.cell(row=row_apps_longer, column=9).value)
        ws_sm.cell(row=row_apps_pub, column=10, value=ws_sm.cell(row=row_apps_longer, column=10).value or 0.0)
        ws_sm.cell(row=row_apps_pub, column=11, value="Commercial: 15, Informational: 4")
        ws_sm.cell(row=row_apps_pub, column=13, value=ws_sm.cell(row=row_apps_longer, column=13).value)
        ws_sm.cell(row=row_apps_pub, column=14, value=ws_sm.cell(row=row_apps_longer, column=14).value)
        ws_sm.cell(row=row_apps_pub, column=15, value=ws_sm.cell(row=row_apps_longer, column=15).value)
        ws_sm.cell(row=row_apps_pub, column=17, value=ws_sm.cell(row=row_apps_longer, column=17).value)
        ws_sm.cell(row=row_apps_pub, column=18, value="Reconciled published page (Legacy Published 52 + Live collection)")
        ws_sm.cell(row=row_apps_pub, column=19, value="Observed in live data sources")
        ws_sm.cell(row=row_apps_pub, column=34, value="Topic mapped")

        ws_sm.cell(row=row_apps_longer, column=2, value="Obsolete planning entry / Alias")
        ws_sm.cell(row=row_apps_longer, column=4, value="Obsolete planning slug (alias of /blog/best-inspection-checklist-apps)")
        ws_sm.cell(row=row_apps_longer, column=18, value="Obsolete planning entry (reconciled to /blog/best-inspection-checklist-apps)")
        ws_sm.cell(row=row_apps_longer, column=19, value="Alias / Never published (live at /blog/best-inspection-checklist-apps)")
        ws_sm.cell(row=row_apps_longer, column=20, value=None)
        ws_sm.cell(row=row_apps_longer, column=21, value=None)
        ws_sm.cell(row=row_apps_longer, column=22, value=None)
        ws_sm.cell(row=row_apps_longer, column=23, value=None)
        ws_sm.cell(row=row_apps_longer, column=33, value="Merged with https://audiment.in/blog/best-inspection-checklist-apps; update any legacy internal links")
        ws_sm.cell(row=row_apps_longer, column=34, value="Obsolete / Merged")

    # Reconcile Features rows
    if row_feat_obsolete and row_feat_pub:
        print(f"Reconciling Row {row_feat_pub} with topic mapping from Row {row_feat_obsolete}...")
        # Col 2: Mapping Type
        ws_sm.cell(row=row_feat_pub, column=2, value="Existing published page")
        # Col 3: Published Title / Planned Topic
        ws_sm.cell(row=row_feat_pub, column=3, value=ws_sm.cell(row=row_feat_obsolete, column=3).value or "Inspection Checklist Software & App Features")
        # Col 4: Content Role / Recommended Type
        ws_sm.cell(row=row_feat_pub, column=4, value="Pillar/commercial guide")
        # Col 5: Cluster ID(s)
        ws_sm.cell(row=row_feat_pub, column=5, value=ws_sm.cell(row=row_feat_obsolete, column=5).value or "C10")
        # Col 6: Canonical Content Topic(s)
        ws_sm.cell(row=row_feat_pub, column=6, value=ws_sm.cell(row=row_feat_obsolete, column=6).value or "Inspection Checklist Software & App Features")
        # Col 7: Pillar(s)
        ws_sm.cell(row=row_feat_pub, column=7, value=ws_sm.cell(row=row_feat_obsolete, column=7).value or "Inspection Management")
        # Col 8: Priority
        ws_sm.cell(row=row_feat_pub, column=8, value=ws_sm.cell(row=row_feat_obsolete, column=8).value or "High")
        # Col 9: Questions Mapped
        ws_sm.cell(row=row_feat_pub, column=9, value=ws_sm.cell(row=row_feat_obsolete, column=9).value or 22)
        # Col 10: Mapped Prompts
        ws_sm.cell(row=row_feat_pub, column=10, value=ws_sm.cell(row=row_feat_obsolete, column=10).value or 22)
        # Col 11: Main Intent Mix
        ws_sm.cell(row=row_feat_pub, column=11, value=ws_sm.cell(row=row_feat_obsolete, column=11).value or "Commercial: 15, Informational: 4, Navigational: 3")
        # Col 12: Representative Prompts
        ws_sm.cell(row=row_feat_pub, column=12, value=ws_sm.cell(row=row_feat_obsolete, column=12).value)
        # Col 13: Primary Existing Hubs
        ws_sm.cell(row=row_feat_pub, column=13, value=ws_sm.cell(row=row_feat_obsolete, column=13).value)
        # Col 14: Supporting Existing Pages
        ws_sm.cell(row=row_feat_pub, column=14, value=ws_sm.cell(row=row_feat_obsolete, column=14).value)
        # Col 15: Planned Cross-Links
        ws_sm.cell(row=row_feat_pub, column=15, value=ws_sm.cell(row=row_feat_obsolete, column=15).value)
        # Col 16: Suggested Incoming Links
        ws_sm.cell(row=row_feat_pub, column=16, value=ws_sm.cell(row=row_feat_obsolete, column=16).value)
        # Col 17: Suggested Outgoing Links
        ws_sm.cell(row=row_feat_pub, column=17, value=ws_sm.cell(row=row_feat_obsolete, column=17).value)
        # Col 18: Source Map Status
        ws_sm.cell(row=row_feat_pub, column=18, value="Reconciled published page (SEO/AEO cluster map + Live collection)")
        # Col 19: Live URL Verification (retain existing observed status)
        ws_sm.cell(row=row_feat_pub, column=19, value="Observed in live data sources")
        # Col 28: Canonical Link URL
        ws_sm.cell(row=row_feat_pub, column=28, value="https://audiment.in/blog/inspection-checklist-software-features")
        # Col 29: Meta Title
        ws_sm.cell(row=row_feat_pub, column=29, value="Inspection Checklist Software & App Features: What to Look For | Audiment")
        # Col 30: Meta Description
        ws_sm.cell(row=row_feat_pub, column=30, value="Compare inspection checklist software features for field teams, including mobile use, GPS, evidence, reporting, workflows, exports, and audit trails.")
        # Col 31: H1 Tag
        ws_sm.cell(row=row_feat_pub, column=31, value="Inspection Checklist Software & App Features: What to Look For")
        # Col 33: Next Action
        ws_sm.cell(row=row_feat_pub, column=33, value="Canonical published entry for topic; monitor indexing and internal link equity")
        # Col 34: Action Status
        ws_sm.cell(row=row_feat_pub, column=34, value="Topic mapped")

        # Mark obsolete planning entry (Row 43)
        print(f"Marking Row {row_feat_obsolete} as obsolete planning URL...")
        ws_sm.cell(row=row_feat_obsolete, column=2, value="Obsolete planning entry / Alias")
        ws_sm.cell(row=row_feat_obsolete, column=4, value="Obsolete planning slug (canonical published at /blog/inspection-checklist-software-features)")
        ws_sm.cell(row=row_feat_obsolete, column=18, value="Obsolete planning entry (reconciled to /blog/inspection-checklist-software-features)")
        ws_sm.cell(row=row_feat_obsolete, column=19, value="Never published / Reconciled to /blog/inspection-checklist-software-features")
        ws_sm.cell(row=row_feat_obsolete, column=20, value=None)
        ws_sm.cell(row=row_feat_obsolete, column=21, value=None)
        ws_sm.cell(row=row_feat_obsolete, column=22, value=None)
        ws_sm.cell(row=row_feat_obsolete, column=23, value=None)
        ws_sm.cell(row=row_feat_obsolete, column=33, value="Never published; canonical live URL is https://audiment.in/blog/inspection-checklist-software-features. Preserved for traceability.")
        ws_sm.cell(row=row_feat_obsolete, column=34, value="Obsolete / Never published")

    # 4. Refresh GSC 28-day metrics and Last Metrics Refresh across all rows in Strategy Map
    refresh_timestamp = datetime.now()
    updated_count = 0
    for r in range(2, ws_sm.max_row + 1):
        norm_u = normalize_url(ws_sm.cell(row=r, column=1).value)
        # Skip obsolete rows from metrics joining
        if r in (row_apps_longer, row_feat_obsolete):
            ws_sm.cell(row=r, column=32, value=refresh_timestamp)
            continue

        if norm_u in url_metrics:
            m = url_metrics[norm_u]
            ws_sm.cell(row=r, column=20, value=m["clicks"])
            ws_sm.cell(row=r, column=21, value=m["impressions"])
            ws_sm.cell(row=r, column=22, value=m["ctr"])
            ws_sm.cell(row=r, column=23, value=m["avg_pos"])
            updated_count += 1

        ws_sm.cell(row=r, column=32, value=refresh_timestamp)

    print(f"Updated 28-day GSC metrics for {updated_count} URLs in Strategy Map.")
    print(f"Last Metrics Refresh timestamp set to: {refresh_timestamp}")

    # 5. Update Internal Linking Sheet
    if 'Internal Linking' in wb.sheetnames:
        ws_il = wb['Internal Linking']
        il_updated_rows = []
        for r in range(2, ws_il.max_row + 1):
            src_url = str(ws_il.cell(row=r, column=3).value or "").strip()
            tgt_url = str(ws_il.cell(row=r, column=5).value or "").strip()
            notes = str(ws_il.cell(row=r, column=7).value or "").strip()

            row_changed = False
            # Check Target URL
            if "inspection-checklist-software-app-features" in tgt_url:
                new_tgt = tgt_url.replace("inspection-checklist-software-app-features", "inspection-checklist-software-features")
                ws_il.cell(row=r, column=5, value=new_tgt)
                row_changed = True
                trace_note = "[Updated target URL from obsolete planning slug /blog/inspection-checklist-software-app-features to canonical published /blog/inspection-checklist-software-features]"
                if trace_note not in notes:
                    ws_il.cell(row=r, column=7, value=f"{notes} {trace_note}".strip())

            # Check Source URL
            if "inspection-checklist-software-app-features" in src_url:
                new_src = src_url.replace("inspection-checklist-software-app-features", "inspection-checklist-software-features")
                ws_il.cell(row=r, column=3, value=new_src)
                row_changed = True
                trace_note = "[Updated source URL from obsolete planning slug /blog/inspection-checklist-software-app-features to canonical published /blog/inspection-checklist-software-features]"
                if trace_note not in str(ws_il.cell(row=r, column=7).value or ""):
                    ws_il.cell(row=r, column=7, value=f"{ws_il.cell(row=r, column=7).value or ''} {trace_note}".strip())

            if row_changed:
                il_updated_rows.append(r)

        print(f"Internal Linking: Updated {len(il_updated_rows)} rows: {il_updated_rows}")

    # 6. Update Content Clusters Sheet
    if 'Content Clusters' in wb.sheetnames:
        ws_cc = wb['Content Clusters']
        cc_updated = []
        for r in range(2, ws_cc.max_row + 1):
            prop_url = str(ws_cc.cell(row=r, column=9).value or "").strip()
            if "inspection-checklist-software-app-features" in prop_url:
                new_prop = prop_url.replace("inspection-checklist-software-app-features", "inspection-checklist-software-features")
                ws_cc.cell(row=r, column=9, value=new_prop)
                cc_updated.append(r)
        print(f"Content Clusters: Updated Proposed URL in {len(cc_updated)} rows: {cc_updated}")

    # 7. Update Question Map Sheet
    if 'Question Map' in wb.sheetnames:
        ws_qm = wb['Question Map']
        qm_updated = []
        for r in range(2, ws_qm.max_row + 1):
            prop_url = str(ws_qm.cell(row=r, column=13).value or "").strip()
            if "inspection-checklist-software-app-features" in prop_url:
                new_prop = prop_url.replace("inspection-checklist-software-app-features", "inspection-checklist-software-features")
                ws_qm.cell(row=r, column=13, value=new_prop)
                qm_updated.append(r)
        print(f"Question Map: Updated Proposed URL in {len(qm_updated)} rows: {qm_updated}")

    wb.save(workbook_path)
    print(f"Successfully saved refreshed workbook to {workbook_path}")
    return True

if __name__ == "__main__":
    targets = [
        "/Users/sanskar/Downloads/Audiment SEO Database (1).xlsx",
        "/Users/sanskar/Downloads/Audiment SEO Database.xlsx"
    ]
    if len(sys.argv) > 1:
        targets = sys.argv[1:]

    for target in targets:
        if os.path.exists(target):
            reconcile_and_refresh_workbook(target)
